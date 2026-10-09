"""Generate size-tiered test files and the test questions that go with them.

    docker compose run --rm testdata                      # all four sizes
    docker compose run --rm testdata --sizes small,medium

Writes eval/sizes/docs/ (one file per type per size) and eval/sizes/cases.json.
Then evaluate them with:

    docker compose run --rm -e EVAL_CONCURRENCY=1 eval \
        --cases eval/sizes/cases.json --label "size tiers"

Everything is synthetic and seeded, so the same files and answers come out
every time. Prose files are an operations manual with one section per depot;
each question targets a fact in one specific section, so the expected answer
and the phrase that must be retrieved are known exactly. Data files are
answered by reference queries, so no answers are hard-coded.

If you copy extra real documents into eval/sizes/docs/ and list them in
eval/sizes/extra_cases.json (same format as cases.json), they are merged in.
"""
import argparse
import csv
import datetime
import io
import json
import os
import random

import config

TIERS_MB = {"small": 0.06, "medium": 1, "large": 10, "xlarge": 45}  # target file sizes
# Prose inside DOCX and PDF stays moderate; images supply the bulk, as in real files.
RICH_SECTIONS = {"small": 30, "medium": 150, "large": 800, "xlarge": 2500}
# Bytes each section and each image add to the finished file, measured per format.
DOCX_BYTES = {"section": 600, "image": 420_000}
PDF_BYTES = {"section": 900, "image": 530_000}

ADJ = ("Amber Birch Cedar Cobalt Copper Coral Crimson Dusty Eastern Emerald Golden Granite Hazel "
       "Indigo Iron Ivory Jade Maple Misty Northern Olive Onyx Pearl Quiet Rustic Sandy Scarlet "
       "Silver Slate Southern Stony Sunny Teal Upper Velvet Western Willow Windy Winter Lower").split()
NOUN = ("Bay Bend Bridge Brook Canal Cove Creek Crossing Dale Dock Falls Field Ford Glen Grove "
        "Harbour Heights Hill Hollow Junction Lake Lane Marsh Meadow Mill Park Pass Pier Point "
        "Reach Ridge Rise Shore Springs Terrace Vale View Wharf Wood Yard").split()
# Names used once per file, for questions that give the depot's name but not its code.
RESERVED = ["Kingfisher Quay", "Saffron Gate", "Orchid Landing", "Pelican Row"]
CITIES = [("Singapore", "Singapore"), ("Johor Bahru", "Malaysia"), ("Penang", "Malaysia"), ("Klang", "Malaysia"),
          ("Batam", "Indonesia"), ("Surabaya", "Indonesia"), ("Jakarta", "Indonesia"), ("Cebu", "Philippines"),
          ("Manila", "Philippines"), ("Da Nang", "Vietnam"), ("Haiphong", "Vietnam"), ("Laem Chabang", "Thailand"),
          ("Bangkok", "Thailand"), ("Phnom Penh", "Cambodia")]
FIRST = "Aisha Ben Chloe Dev Elena Farid Grace Hiro Ivan Jia Kofi Lena Marco Nadia Omar Priya Rosa Sam Tara Wei".split()
LAST = "Ng Lim Wong Kumar Santos Ibrahim Chen Tanaka Novak Singh Mensah Fischer Rossi Hassan Lee Reyes Park Goh Shah Aziz".split()
ASSEMBLY = ["the staff car park", "the grass area beside the main gate", "the container yard entrance",
            "the lorry holding bay", "the visitor car park", "the forecourt of loading bay 12",
            "the field behind Warehouse 2", "the north car park by Gate 3"]
COLD = ["2 to 8 degrees Celsius", "0 to 4 degrees Celsius", "15 to 25 degrees Celsius", "minus 20 to minus 15 degrees Celsius"]

INTRO = ("Harbourline Freight Depot Operations Manual. This manual describes every depot in the Harbourline "
         "Freight network. Each depot has its own section, identified by a depot code, which names the depot "
         "manager and sets out the canteen hours, staff shuttle, parking, fire assembly point, dock doors, "
         "same-day dispatch cut-off, cold storage range and security desk extension for that depot. "
         "This edition covers {n:,} depots.")
SUMMARY_POINTS = [
    "The document is an operations manual for Harbourline Freight depots",
    "Each depot section names the depot manager",
    "Sections give canteen hours and staff shuttle details",
    "Sections give parking and fire assembly point information",
    "Sections give dock door numbers and the same-day dispatch cut-off",
]


def _clock(rng, lo, hi):
    return f"{rng.randint(lo, hi):02d}:{rng.choice(['00', '15', '30', '45'])}"


def make_depot(i, rng, name=None):
    """One manual section: its text, and the facts questions can be asked about."""
    code = f"HD-{i:05d}"
    name = name or f"{rng.choice(ADJ)} {rng.choice(NOUN)}"
    city, country = rng.choice(CITIES)
    d = {"code": code, "name": name, "manager": f"{rng.choice(FIRST)} {rng.choice(LAST)}",
         "open": _clock(rng, 5, 7), "close": _clock(rng, 17, 21), "shuttle": _clock(rng, 5, 6),
         "spaces": rng.randint(20, 400), "docks": rng.randint(4, 60), "cutoff": _clock(rng, 13, 18),
         "forklifts": rng.randint(2, 45), "ext": rng.randint(2000, 8999)}
    d["heading"] = f"{i}. {name} depot ({code})"
    d["ev_manager"] = f"({code}) in {city}, {country} is managed by {d['manager']}"
    d["ev_close"] = f"The canteen at {code} is open from {d['open']} to {d['close']}"
    d["ev_spaces"] = f"There are {d['spaces']} parking spaces at {code}"
    d["ev_cutoff"] = f"same-day dispatch cut-off at {code} is {d['cutoff']}"
    d["ev_ext"] = f"security desk at {code} can be reached on extension {d['ext']}"
    d["paragraphs"] = [
        f"The {name} depot {d['ev_manager']}, who is the first point of contact for access passes, building "
        f"issues and local safety arrangements. New joiners receive a site induction on their first morning, "
        f"before they are allowed onto the warehouse floor.",
        f"{d['ev_close']} on working days and accepts the staff meal card. A free shuttle bus leaves the "
        f"nearest station at {d['shuttle']} for the early shift and again eight hours later for the late shift.",
        f"{d['ev_spaces']}, allocated on a first-come basis; those nearest the entrance are reserved for "
        f"visitors and for colleagues with reduced mobility. If the fire alarm sounds, leave by the nearest "
        f"exit and go to the assembly point, which is {rng.choice(ASSEMBLY)}.",
        f"The depot has {d['docks']} dock doors and {d['forklifts']} forklift trucks. The {d['ev_cutoff']}; "
        f"orders released to the floor after that time are dispatched on the next working day.",
        f"Cold storage is held at {rng.choice(COLD)}. The {d['ev_ext']} at any hour, and safety shoes and a "
        f"high-visibility vest must be worn at all times in the warehouse.",
    ]
    return d


def make_manual(n_sections, seed):
    rng = random.Random(seed)
    special = max(1, int(n_sections * 0.72))  # the section that gets a one-off name
    name = RESERVED[seed % len(RESERVED)]
    return [make_depot(i, rng, name if i == special else None) for i in range(1, n_sections + 1)], special


def prose_cases(depots, special, prefix):
    n = len(depots)
    at = lambda share: depots[min(n - 1, int(n * share))]
    a, b, c, s, f = at(0.05), at(0.50), at(0.95), depots[special - 1], at(0.33)
    label = lambda d: f"the {d['name']} depot ({d['code']})"
    cases = [
        {"type": "answerable", "question": f"What time does the canteen close at {label(a)}?",
         "expected": f"{a['close']} (it is open from {a['open']} to {a['close']}).", "evidence": a["ev_close"]},
        {"type": "answerable", "question": f"What is the same-day dispatch cut-off at {label(b)}?",
         "expected": f"{b['cutoff']}.", "evidence": b["ev_cutoff"]},
        {"type": "answerable", "question": f"Which phone extension reaches the security desk at {label(c)}?",
         "expected": f"Extension {c['ext']}.", "evidence": c["ev_ext"]},
        # Name only, no code: the name appears in one paragraph, the answer in a later one.
        {"type": "answerable", "question": f"How many parking spaces are there at the {s['name']} depot?",
         "expected": f"{s['spaces']} spaces.", "evidence": s["ev_spaces"]},
        {"type": "followup", "previous": {"question": f"Who manages {label(f)}?", "answer": f"{f['manager']}."},
         "question": "And when does its canteen open?", "expected": f"{f['open']}.", "evidence": f["ev_close"]},
        {"type": "unanswerable", "question": f"What is the annual rent paid for {label(b)}?"},
    ]
    return [{"id": f"{prefix}-{i:02d}", **case} for i, case in enumerate(cases, start=1)]


# --- prose writers -----------------------------------------------------------

def write_txt(path, depots):
    with open(path, "w") as f:
        f.write(INTRO.format(n=len(depots)) + "\n")
        for d in depots:
            f.write(f"\n\n{d['heading']}\n\n" + "\n\n".join(d["paragraphs"]) + "\n")


def write_html(path, depots):
    with open(path, "w") as f:
        f.write('<!DOCTYPE html>\n<html lang="en"><head><meta charset="utf-8"><title>Depot Operations Manual</title>\n'
                "<style>body{font-family:sans-serif} section{margin:2em 0} nav a{margin-right:1em}</style>\n"
                '<script>window.analytics = {track: function (e) { console.log("event", e); }};</script>\n</head><body>\n'
                '<nav><a href="/">Home</a><a href="/depots">Depots</a><a href="/safety">Safety</a></nav>\n'
                f"<h1>Depot Operations Manual</h1>\n<p>{INTRO.format(n=len(depots))}</p>\n")
        for d in depots:
            f.write(f'<section id="{d["code"]}"><h2>{d["heading"]}</h2>\n'
                    + "\n".join(f"<p>{p}</p>" for p in d["paragraphs"]) + "\n</section>\n")
        f.write('<footer><p>Harbourline Freight internal document.</p></footer>\n<script>analytics.track("view");</script>\n</body></html>\n')


def _images(n, seed):
    """n distinct JPEGs of random noise (incompressible, so they give files real weight)."""
    import numpy as np
    from PIL import Image
    rng = np.random.default_rng(seed)
    for _ in range(n):
        buf = io.BytesIO()
        Image.fromarray(rng.integers(0, 255, (560, 900, 3), dtype="uint8")).save(buf, "JPEG", quality=88)
        buf.seek(0)
        yield buf


def _image_plan(depots, target_bytes, weights):
    """How many images to add, and after which sections, to bring a file up to size."""
    n = max(0, int((target_bytes - len(depots) * weights["section"]) / weights["image"]))
    every = max(1, len(depots) // n) if n else 0
    return n, every


def write_docx(path, depots, target_bytes, seed):
    import docx
    from docx.shared import Cm
    n, every = _image_plan(depots, target_bytes, DOCX_BYTES)
    images = _images(n, seed)
    doc = docx.Document()
    doc.add_heading("Depot Operations Manual", 0)
    doc.add_paragraph(INTRO.format(n=len(depots)))
    for i, d in enumerate(depots, start=1):
        doc.add_heading(d["heading"], 2)
        for p in d["paragraphs"]:
            doc.add_paragraph(p)
        if every and i % every == 0 and (img := next(images, None)):
            doc.add_picture(img, width=Cm(15))
            doc.add_paragraph(f"Figure: site photograph for {d['code']}.")
    doc.save(path)


def write_pdf(path, depots, target_bytes, seed):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
    n, every = _image_plan(depots, target_bytes, PDF_BYTES)
    images = _images(n, seed)
    styles = getSampleStyleSheet()
    flow = [Paragraph("Depot Operations Manual", styles["Title"]),
            Paragraph(INTRO.format(n=len(depots)), styles["Normal"]), Spacer(1, 14)]
    for i, d in enumerate(depots, start=1):
        flow.append(Paragraph(d["heading"], styles["Heading2"]))
        for p in d["paragraphs"]:
            flow += [Paragraph(p, styles["Normal"]), Spacer(1, 8)]
        if every and i % every == 0 and (img := next(images, None)):
            flow += [Image(img, width=15 * cm, height=9.3 * cm),
                     Paragraph(f"Figure: site photograph for {d['code']}.", styles["Italic"]), Spacer(1, 10)]
    SimpleDocTemplate(path, pagesize=A4, title="Depot Operations Manual").build(flow)


# --- data writers ------------------------------------------------------------

PRODUCTS = {"Pallet Wrap": 18.5, "Shipping Label": 4.2, "Carton Box": 2.75, "Tape Roll": 3.6, "Bubble Mailer": 0.95}


def write_csv(path, target_bytes, seed):
    rng = random.Random(seed)
    start = datetime.date(2025, 1, 1)
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Order ID", "Order Date", "Region", "Product", "Units", "Unit Price", "Channel", "Status"])
        i = 0
        while f.tell() < target_bytes:
            i += 1
            p = rng.choices(list(PRODUCTS), weights=[2, 5, 6, 4, 3])[0]
            w.writerow([f"ORD-{i:07d}", (start + datetime.timedelta(days=rng.randrange(365))).isoformat(),
                        rng.choices(["North", "South", "East", "West"], weights=[3, 2, 4, 2])[0], p,
                        rng.randint(1, 50), round(PRODUCTS[p] * rng.choice([1, 1, 1, 0.9, 1.1]), 2),
                        rng.choices(["Online", "Reseller", "Direct"], weights=[5, 3, 2])[0],
                        rng.choices(["Delivered", "Returned", "Cancelled"], weights=[90, 6, 4])[0]])


CSV_CASES = [
    ("answerable", "How many orders are in the file?", "The total number of orders.", "SELECT COUNT(*) FROM data"),
    ("answerable", "Which region has the most orders?", "The region with the highest order count.",
     "SELECT region FROM data GROUP BY region ORDER BY COUNT(*) DESC LIMIT 1"),
    ("answerable", "What is the total revenue, counting revenue as units times unit price?", "Total revenue across all orders.",
     "SELECT ROUND(SUM(units * unit_price), 2) FROM data"),
    ("answerable", "How many orders were returned?", "The number of orders with status Returned.",
     "SELECT COUNT(*) FROM data WHERE status = 'Returned'"),
    ("answerable", "What was the average order value in December 2025?", "The mean of units times unit price for December 2025 orders.",
     "SELECT ROUND(AVG(units * unit_price), 2) FROM data WHERE order_date LIKE '2025-12-%'"),
    ("unanswerable", "What was the profit margin on Carton Box?", None, None),
]

XLSX_ROWS = {"small": 1_100, "medium": 19_000, "large": 190_000, "xlarge": 860_000}  # about 55 bytes a row once zipped
DEPOTS = [("Tuas", "Singapore"), ("Jurong", "Singapore"), ("Johor", "Malaysia"), ("Penang", "Malaysia"),
          ("Klang", "Malaysia"), ("Batam", "Indonesia"), ("Surabaya", "Indonesia"), ("Cebu", "Philippines")]


def write_xlsx(path, rows, seed):
    from openpyxl import Workbook
    rng = random.Random(seed)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Shipments")
    ws.append(["Shipment ID", "Ship Date", "Depot", "Mode", "Service", "Weight Kg", "Cost", "Tracking Ref"])
    start = datetime.date(2025, 1, 1)
    for i in range(1, rows + 1):
        weight = round(rng.uniform(0.5, 900), 1)
        ws.append([f"S{i:07d}", (start + datetime.timedelta(days=rng.randrange(365))).isoformat(),
                   rng.choices(DEPOTS, weights=[6, 5, 4, 3, 3, 2, 2, 1])[0][0],
                   rng.choices(["Road", "Sea", "Air"], weights=[6, 3, 1])[0],
                   rng.choices(["Standard", "Express"], weights=[4, 1])[0],
                   weight, round(weight * rng.uniform(0.4, 2.2), 2), f"{rng.getrandbits(64):016x}"])
    ws2 = wb.create_sheet("Depots")
    ws2.append(["Depot", "Country", "Dock Doors"])
    for k, (depot, country) in enumerate(DEPOTS):
        ws2.append([depot, country, 12 + 6 * k])
    wb.save(path)


XLSX_CASES = [
    ("answerable", "How many shipments are there?", "The total number of shipments.", "SELECT COUNT(*) FROM shipments"),
    ("answerable", "Which depot handled the most shipments?", "The depot with the highest shipment count.",
     "SELECT depot FROM shipments GROUP BY depot ORDER BY COUNT(*) DESC LIMIT 1"),
    ("answerable", "What is the total weight shipped by road, in kilograms?", "Total weight of Road shipments.",
     "SELECT ROUND(SUM(weight_kg), 1) FROM shipments WHERE mode = 'Road'"),
    ("answerable", "What is the average cost of an Express shipment?", "Mean cost of Express shipments.",
     "SELECT ROUND(AVG(cost), 2) FROM shipments WHERE service = 'Express'"),
    ("answerable", "How many shipments went through depots in Malaysia?", "The number of shipments from depots whose country is Malaysia.",
     "SELECT COUNT(*) FROM shipments s JOIN depots d ON d.depot = s.depot WHERE d.country = 'Malaysia'"),
    ("unanswerable", "Which driver delivered the most shipments?", None, None),
]

CATEGORIES = {"Billing": 6, "Delivery delay": 20, "Damaged goods": 34, "Account access": 3, "API": 12}


def write_json(path, target_bytes, seed):
    rng = random.Random(seed)
    start = datetime.date(2026, 1, 1)
    with open(path, "w") as f:
        f.write("[\n")
        i = 0
        while f.tell() < target_bytes:
            i += 1
            cat = rng.choices(list(CATEGORIES), weights=[4, 7, 3, 3, 2])[0]
            pr = rng.choices(["Low", "Medium", "High"], weights=[4, 4, 2])[0]
            rec = {"id": f"T-{i:07d}", "created": (start + datetime.timedelta(days=rng.randrange(180))).isoformat(),
                   "customer": {"name": f"{rng.choice(LAST)} {rng.choice(['Trading', 'Retail', 'Online', 'Foods', 'Home'])}",
                                "tier": rng.choices(["Bronze", "Silver", "Gold"], weights=[5, 3, 2])[0]},
                   "category": cat, "priority": pr,
                   "resolution_hours": round(CATEGORIES[cat] * rng.uniform(0.4, 1.8) * {"Low": 1.2, "Medium": 1.0, "High": 0.6}[pr], 1),
                   "satisfied": rng.random() < (0.55 if cat == "Damaged goods" else 0.82),
                   "tags": rng.sample(["refund", "urgent", "repeat", "escalated", "vip"], k=rng.randint(0, 2))}
            f.write(("" if i == 1 else ",\n") + json.dumps(rec))
        f.write("\n]\n")


JSON_CASES = [
    ("answerable", "How many tickets are there?", "The total number of tickets.", "SELECT COUNT(*) FROM data"),
    ("answerable", "What is the most common ticket category?", "The category with the most tickets.",
     "SELECT category FROM data GROUP BY category ORDER BY COUNT(*) DESC LIMIT 1"),
    ("answerable", "What is the average resolution time for high-priority tickets?", "Mean resolution hours for High priority tickets.",
     "SELECT ROUND(AVG(resolution_hours), 2) FROM data WHERE priority = 'High'"),
    ("answerable", "How many tickets came from Gold-tier customers?", "The number of tickets from Gold tier customers.",
     "SELECT COUNT(*) FROM data WHERE customer_tier = 'Gold'"),
    ("answerable", "Which category takes longest to resolve on average?", "The category with the highest mean resolution hours.",
     "SELECT category FROM data GROUP BY category ORDER BY AVG(resolution_hours) DESC LIMIT 1"),
    ("unanswerable", "Which support agent closed the most tickets?", None, None),
]
DATA_POINTS = {
    "csv": ["The file contains sales orders", "Orders are dated in 2025", "There are four regions",
            "There are five products", "Each order has a status such as delivered, returned or cancelled"],
    "xlsx": ["The workbook contains shipment records", "There is a second sheet listing depots",
             "Shipments travel by road, sea or air", "Each shipment has a weight and a cost"],
    "json": ["The file contains support tickets", "Tickets have a category such as billing or delivery delay",
             "Tickets have a priority level", "Resolution time in hours is recorded", "Customers have a tier"],
}


def data_cases(spec, prefix):
    out = []
    for i, (kind, question, expected, sql) in enumerate(spec, start=1):
        case = {"id": f"{prefix}-{i:02d}", "type": kind, "question": question}
        if expected:
            case.update(expected=expected, gold_sql=sql)
        out.append(case)
    return out


# --- main --------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Generate size-tiered test files.")
    ap.add_argument("--sizes", default=",".join(TIERS_MB), help="Comma-separated: small,medium,large,xlarge")
    ap.add_argument("--out", default=os.path.join(config.EVAL_DIR, "sizes"))
    args = ap.parse_args()
    docs_dir = os.path.join(args.out, "docs")
    os.makedirs(docs_dir, exist_ok=True)

    documents = []
    for tier in args.sizes.split(","):
        target = int(TIERS_MB[tier] * 1024 * 1024)
        seed = 1000 + list(TIERS_MB).index(tier)

        def add(name, writer, cases, points):
            path = os.path.join(docs_dir, name)
            writer(path)
            mb = os.path.getsize(path) / 1024 / 1024
            print(f"  {name:<22}{mb:>8.2f} MB", flush=True)
            documents.append({"file": name, "size": tier, "summary_points": points, "cases": cases})

        print(f"{tier} (target about {TIERS_MB[tier]} MB)", flush=True)
        # Plain prose: the text itself is sized to the target (about 1,300 bytes a section).
        manual, special = make_manual(max(8, target // 1300), seed)
        add(f"manual_{tier}.txt", lambda p: write_txt(p, manual), prose_cases(manual, special, f"txt-{tier}"), SUMMARY_POINTS)
        html_manual, special = make_manual(max(8, target // 1400), seed + 10)
        add(f"manual_{tier}.html", lambda p: write_html(p, html_manual),
            prose_cases(html_manual, special, f"html-{tier}"), SUMMARY_POINTS)
        # Rich documents: moderate text, images make up the size.
        rich, special = make_manual(RICH_SECTIONS[tier], seed + 20)
        add(f"manual_{tier}.docx", lambda p: write_docx(p, rich, target, seed), prose_cases(rich, special, f"docx-{tier}"), SUMMARY_POINTS)
        rich_pdf, special = make_manual(RICH_SECTIONS[tier], seed + 30)
        add(f"manual_{tier}.pdf", lambda p: write_pdf(p, rich_pdf, target, seed + 1),
            prose_cases(rich_pdf, special, f"pdf-{tier}"), SUMMARY_POINTS)
        # Data files.
        add(f"orders_{tier}.csv", lambda p: write_csv(p, target, seed), data_cases(CSV_CASES, f"csv-{tier}"), DATA_POINTS["csv"])
        add(f"shipments_{tier}.xlsx", lambda p: write_xlsx(p, XLSX_ROWS[tier], seed),
            data_cases(XLSX_CASES, f"xlsx-{tier}"), DATA_POINTS["xlsx"])
        add(f"tickets_{tier}.json", lambda p: write_json(p, target, seed), data_cases(JSON_CASES, f"json-{tier}"), DATA_POINTS["json"])

    # Real documents the user has added alongside the generated ones.
    extra_path = os.path.join(args.out, "extra_cases.json")
    if os.path.exists(extra_path):
        with open(extra_path) as f:
            for doc in json.load(f)["documents"]:
                if os.path.exists(os.path.join(docs_dir, doc["file"])):
                    documents.append(doc)
                    print(f"Added {doc['file']} from extra_cases.json")
                else:
                    print(f"Skipped {doc['file']}: listed in extra_cases.json but not found in {docs_dir}")

    with open(os.path.join(args.out, "cases.json"), "w") as f:
        json.dump({"note": "Generated by app/make_testdata.py; rerun it rather than editing by hand.",
                   "documents": documents}, f, indent=2)
    print(f"Wrote {len(documents)} documents and {sum(len(d['cases']) for d in documents)} questions to {args.out}")


if __name__ == "__main__":
    main()
